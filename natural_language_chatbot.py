import os
import json
import sqlite3
import pandas as pd
import numpy as np
from datetime import datetime
import tkinter as tk
from tkinter import ttk, scrolledtext, filedialog, messagebox
import threading
import queue
import time
import logging
from pathlib import Path
import re

# AI/ML imports
import torch
from transformers import (
    AutoTokenizer, AutoModelForCausalLM, 
    BitsAndBytesConfig, pipeline
)
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity

# Web search imports
import requests
from bs4 import BeautifulSoup
from googlesearch import search
import urllib.parse

# Document processing
import docx
from PyPDF2 import PdfReader
from openpyxl import load_workbook
import easyocr

class NaturalLanguageSalesChatBot:
    """
    AI Sales ChatBot with natural language generation using local LLM
    Priority: Local Database → Google Search → Fallback Response
    """
    
    def __init__(self):
        self.setup_logging()
        self.setup_database()
        self.initialize_ai_models()
        self.conversation_context = []
        self.setup_gui()
        
    def setup_logging(self):
        """Setup logging configuration"""
        logging.basicConfig(
            level=logging.INFO,
            format='%(asctime)s - %(levelname)s - %(message)s',
            handlers=[
                logging.FileHandler('chatbot.log'),
                logging.StreamHandler()
            ]
        )
        self.logger = logging.getLogger(__name__)
        
    def setup_database(self):
        """Initialize SQLite database"""
        self.conn = sqlite3.connect('chatbot_data.db', check_same_thread=False)
        cursor = self.conn.cursor()
        
        # Enhanced database schema
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                description TEXT,
                category TEXT,
                price REAL,
                features TEXT,
                specifications TEXT,
                availability TEXT,
                source_file TEXT,
                embedding BLOB,
                created_at TEXT,
                updated_at TEXT
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS knowledge_base (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                topic TEXT NOT NULL,
                content TEXT NOT NULL,
                keywords TEXT,
                source TEXT,
                embedding BLOB,
                created_at TEXT
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT,
                user_input TEXT,
                bot_response TEXT,
                intent TEXT,
                confidence REAL,
                data_source TEXT,
                response_time REAL
            )
        ''')
        
        self.conn.commit()
        
    def initialize_ai_models(self):
        """Initialize AI models for natural language generation"""
        self.logger.info("Initializing AI models...")
        
        try:
            # Device setup for heterogeneous GPUs
            self.setup_devices()
            
            # Load LLM for natural language generation
            self.load_language_model()
            
            # Load embedding model for semantic search
            self.load_embedding_model()
            
            # Load OCR model
            self.load_ocr_model()
            
            self.logger.info("All AI models loaded successfully!")
            
        except Exception as e:
            self.logger.error(f"Error initializing AI models: {e}")
            self.models = {}
            
    def setup_devices(self):
        """Setup device configuration for optimal GPU usage"""
        if torch.cuda.is_available():
            gpu_count = torch.cuda.device_count()
            self.devices = []
            
            for i in range(gpu_count):
                props = torch.cuda.get_device_properties(i)
                self.devices.append({
                    'id': i,
                    'name': props.name,
                    'memory_gb': props.total_memory / (1024**3),
                    'device': f'cuda:{i}'
                })
                
            # Sort by memory (largest first for LLM)
            self.devices.sort(key=lambda x: x['memory_gb'], reverse=True)
            
            self.primary_device = self.devices[0]['device']  # For LLM
            self.secondary_device = self.devices[1]['device'] if len(self.devices) > 1 else self.primary_device
            
            print(f"Primary GPU (LLM): {self.devices[0]['name']} - {self.devices[0]['memory_gb']:.1f}GB")
            if len(self.devices) > 1:
                print(f"Secondary GPU (Embeddings): {self.devices[1]['name']} - {self.devices[1]['memory_gb']:.1f}GB")
        else:
            self.primary_device = 'cpu'
            self.secondary_device = 'cpu'
            print("CUDA not available. Using CPU mode.")
            
    def load_language_model(self):
        """Load language model for natural text generation"""
        try:
            # Model selection based on available VRAM
            primary_memory = self.devices[0]['memory_gb'] if torch.cuda.is_available() else 8
            
            if primary_memory >= 20:  # RTX 4090 or similar
                model_name = "meta-llama/Llama-2-7b-chat-hf"
                use_quantization = False
            elif primary_memory >= 12:  # RTX 4070Ti Super or similar
                model_name = "meta-llama/Llama-2-7b-chat-hf"
                use_quantization = True
            else:
                model_name = "microsoft/DialoGPT-medium"
                use_quantization = True
                
            print(f"Loading language model: {model_name}")
            
            # Configure quantization if needed
            if use_quantization:
                quantization_config = BitsAndBytesConfig(
                    load_in_8bit=True,
                    llm_int8_threshold=6.0,
                )
            else:
                quantization_config = None
                
            # Load tokenizer
            self.tokenizer = AutoTokenizer.from_pretrained(model_name)
            if self.tokenizer.pad_token is None:
                self.tokenizer.pad_token = self.tokenizer.eos_token
                
            # Load model
            self.llm_model = AutoModelForCausalLM.from_pretrained(
                model_name,
                quantization_config=quantization_config,
                device_map={"": self.primary_device},
                torch_dtype=torch.float16 if torch.cuda.is_available() else torch.float32,
                trust_remote_code=True
            )
            
            # Create text generation pipeline
            self.text_generator = pipeline(
                "text-generation",
                model=self.llm_model,
                tokenizer=self.tokenizer,
                device=0 if torch.cuda.is_available() else -1,
                torch_dtype=torch.float16 if torch.cuda.is_available() else torch.float32
            )
            
            print("✅ Language model loaded successfully")
            
        except Exception as e:
            print(f"❌ Error loading language model: {e}")
            self.load_fallback_model()
            
    def load_fallback_model(self):
        """Load a simpler fallback model"""
        try:
            print("Loading fallback model: microsoft/DialoGPT-medium")
            
            self.tokenizer = AutoTokenizer.from_pretrained("microsoft/DialoGPT-medium")
            self.llm_model = AutoModelForCausalLM.from_pretrained("microsoft/DialoGPT-medium")
            
            if torch.cuda.is_available():
                self.llm_model = self.llm_model.to(self.primary_device)
                
            self.text_generator = pipeline(
                "text-generation",
                model=self.llm_model,
                tokenizer=self.tokenizer,
                device=0 if torch.cuda.is_available() else -1
            )
            
            print("✅ Fallback model loaded successfully")
            
        except Exception as e:
            print(f"❌ Error loading fallback model: {e}")
            self.text_generator = None
            
    def load_embedding_model(self):
        """Load embedding model for semantic search"""
        try:
            model_name = 'sentence-transformers/all-MiniLM-L6-v2'
            self.embedding_model = SentenceTransformer(model_name, device=self.secondary_device)
            print("✅ Embedding model loaded successfully")
        except Exception as e:
            print(f"❌ Error loading embedding model: {e}")
            self.embedding_model = None
            
    def load_ocr_model(self):
        """Load OCR model for image processing"""
        try:
            self.ocr_reader = easyocr.Reader(['en'])
            print("✅ OCR model loaded successfully")
        except Exception as e:
            print(f"❌ Error loading OCR model: {e}")
            self.ocr_reader = None
            
    def generate_natural_response(self, user_input, context_data=None, data_source="unknown"):
        """Generate natural language response using LLM"""
        try:
            # Build context-aware prompt
            prompt = self.build_sales_prompt(user_input, context_data, data_source)
            
            if not self.text_generator:
                return "I apologize, but the AI text generation system is not available. Please try again later."
                
            # Generate response
            with torch.no_grad():
                generated = self.text_generator(
                    prompt,
                    max_length=len(prompt.split()) + 150,  # Dynamic length
                    min_length=len(prompt.split()) + 20,
                    temperature=0.7,
                    do_sample=True,
                    pad_token_id=self.tokenizer.eos_token_id,
                    num_return_sequences=1,
                    repetition_penalty=1.1
                )
                
            # Extract generated text
            full_response = generated[0]['generated_text']
            
            # Extract only the new part (after the prompt)
            response = full_response[len(prompt):].strip()
            
            # Clean up the response
            response = self.clean_generated_response(response)
            
            return response
            
        except Exception as e:
            self.logger.error(f"Error generating response: {e}")
            return "I apologize, but I'm having trouble generating a response right now. Could you please rephrase your question?"
            
    def build_sales_prompt(self, user_input, context_data=None, data_source="unknown"):
        """Build an optimized prompt for sales conversations"""
        
        # Base system prompt
        prompt = """You are a professional AI sales assistant for a company. You help customers find products, answer questions, and provide excellent customer service.

Instructions:
- Be helpful, friendly, and professional
- Provide specific product information when available
- If you don't know something, be honest about it
- Focus on solving the customer's needs
- Keep responses concise but informative

"""
        
        # Add conversation context
        if self.conversation_context:
            prompt += "Previous conversation:\n"
            for turn in self.conversation_context[-3:]:  # Last 3 turns
                prompt += f"Customer: {turn['user']}\nAssistant: {turn['bot']}\n"
            prompt += "\n"
            
        # Add relevant data context
        if context_data:
            if data_source == "database":
                prompt += "Relevant products from our catalog:\n"
                for item in context_data[:3]:  # Top 3 results
                    prompt += f"- {item.get('name', 'Product')}: {item.get('description', 'No description')[:100]}...\n"
                    if item.get('price'):
                        prompt += f"  Price: ${item['price']}\n"
                prompt += "\n"
                
            elif data_source == "web_search":
                prompt += "Additional information found online:\n"
                for item in context_data[:2]:  # Top 2 web results
                    prompt += f"- {item[:150]}...\n"
                prompt += "\n"
                
        # Current user question
        prompt += f"Current customer question: {user_input}\n\n"
        prompt += "Assistant response:"
        
        return prompt
        
    def clean_generated_response(self, response):
        """Clean up generated response"""
        # Remove common artifacts
        response = re.sub(r'<[^>]+>', '', response)  # Remove HTML tags
        response = re.sub(r'\[.*?\]', '', response)  # Remove brackets
        response = re.sub(r'\n+', ' ', response)     # Replace newlines with spaces
        response = re.sub(r'\s+', ' ', response)     # Collapse multiple spaces
        
        # Remove incomplete sentences at the end
        sentences = response.split('.')
        if len(sentences) > 1 and len(sentences[-1].strip()) < 10:
            response = '.'.join(sentences[:-1]) + '.'
            
        return response.strip()
        
    def search_local_database(self, user_input, similarity_threshold=0.7):
        """Search local database using semantic similarity"""
        try:
            if not self.embedding_model:
                return []
                
            # Get query embedding
            query_embedding = self.embedding_model.encode([user_input])
            
            # Search products
            cursor = self.conn.cursor()
            cursor.execute("""
                SELECT name, description, category, price, features, specifications, availability 
                FROM products
            """)
            products = cursor.fetchall()
            
            if not products:
                return []
                
            # Calculate similarities
            best_matches = []
            
            for product in products:
                # Combine product text for embedding
                product_text = f"{product[0]} {product[1]} {product[4]} {product[5]}"
                product_embedding = self.embedding_model.encode([product_text])
                
                # Calculate similarity
                similarity = cosine_similarity(query_embedding, product_embedding)[0][0]
                
                if similarity > similarity_threshold:
                    best_matches.append({
                        'name': product[0],
                        'description': product[1],
                        'category': product[2],
                        'price': product[3],
                        'features': product[4],
                        'specifications': product[5],
                        'availability': product[6],
                        'similarity': similarity
                    })
                    
            # Sort by similarity
            best_matches.sort(key=lambda x: x['similarity'], reverse=True)
            
            return best_matches[:5]  # Top 5 matches
            
        except Exception as e:
            self.logger.error(f"Error searching local database: {e}")
            return []
            
    def search_google(self, query, num_results=3):
        """Search Google for additional information"""
        try:
            # Enhanced search query for better results
            search_query = f"{query} product information specifications price"
            
            search_results = []
            
            # Use googlesearch library
            for url in search(search_query, num=num_results, stop=num_results, pause=1):
                try:
                    # Fetch page content
                    headers = {
                        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
                    }
                    
                    response = requests.get(url, headers=headers, timeout=5)
                    if response.status_code == 200:
                        soup = BeautifulSoup(response.content, 'html.parser')
                        
                        # Extract text content
                        text_content = soup.get_text()
                        
                        # Clean and truncate
                        clean_text = re.sub(r'\s+', ' ', text_content).strip()
                        
                        if len(clean_text) > 200:
                            search_results.append(clean_text[:500])
                            
                except Exception as e:
                    continue
                    
            return search_results
            
        except Exception as e:
            self.logger.error(f"Error searching Google: {e}")
            return []
            
    def process_user_message(self, user_input):
        """Main processing pipeline with natural language generation"""
        start_time = time.time()
        
        try:
            self.logger.info(f"Processing user input: {user_input}")
            
            # Step 1: Search local database first
            local_results = self.search_local_database(user_input)
            
            if local_results:
                # Generate response based on local data
                response = self.generate_natural_response(
                    user_input, 
                    context_data=local_results,
                    data_source="database"
                )
                data_source = "local_database"
                self.logger.info("Response generated from local database")
                
            else:
                # Step 2: Fallback to Google search
                self.logger.info("No local results found, searching Google...")
                web_results = self.search_google(user_input)
                
                if web_results:
                    response = self.generate_natural_response(
                        user_input,
                        context_data=web_results,
                        data_source="web_search"
                    )
                    data_source = "google_search"
                    self.logger.info("Response generated from Google search")
                    
                else:
                    # Step 3: Generate general helpful response
                    response = self.generate_natural_response(user_input)
                    data_source = "general_knowledge"
                    self.logger.info("Generated general response")
                    
            # Update conversation context
            self.conversation_context.append({
                'user': user_input,
                'bot': response,
                'timestamp': datetime.now().isoformat()
            })
            
            # Keep only last 10 turns
            if len(self.conversation_context) > 10:
                self.conversation_context = self.conversation_context[-10:]
                
            # Store conversation in database
            processing_time = time.time() - start_time
            self.store_conversation(user_input, response, data_source, processing_time)
            
            return {
                'response': response,
                'data_source': data_source,
                'local_results_count': len(local_results),
                'processing_time': processing_time
            }
            
        except Exception as e:
            error_msg = f"I apologize, but I encountered an error processing your request. Please try again."
            self.logger.error(f"Error processing user message: {e}")
            return {
                'response': error_msg,
                'data_source': 'error',
                'local_results_count': 0,
                'processing_time': time.time() - start_time
            }
            
    def store_conversation(self, user_input, response, data_source, processing_time):
        """Store conversation in database"""
        try:
            cursor = self.conn.cursor()
            cursor.execute('''
                INSERT INTO conversations 
                (timestamp, user_input, bot_response, data_source, response_time)
                VALUES (?, ?, ?, ?, ?)
            ''', (
                datetime.now().isoformat(),
                user_input,
                response,
                data_source,
                processing_time
            ))
            self.conn.commit()
        except Exception as e:
            self.logger.error(f"Error storing conversation: {e}")
            
    def setup_gui(self):
        """Setup the GUI interface"""
        self.root = tk.Tk()
        self.root.title("Natural Language AI Sales ChatBot")
        self.root.geometry("1400x900")
        
        # Create notebook for tabs
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Chat tab
        self.setup_chat_tab()
        
        # Database management tab
        self.setup_database_tab()
        
        # Training tab
        self.setup_training_tab()
        
        # Analytics tab
        self.setup_analytics_tab()
        
    def setup_chat_tab(self):
        """Setup the main chat interface"""
        chat_frame = ttk.Frame(self.notebook)
        self.notebook.add(chat_frame, text="💬 Chat")
        
        # Chat display with styling
        self.chat_display = scrolledtext.ScrolledText(
            chat_frame, 
            height=30, 
            state=tk.DISABLED,
            font=('Consolas', 10),
            wrap=tk.WORD
        )
        self.chat_display.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # Configure text tags for styling
        self.chat_display.tag_configure("user", foreground="blue")
        self.chat_display.tag_configure("bot", foreground="green")
        self.chat_display.tag_configure("system", foreground="gray")
        
        # Input frame
        input_frame = ttk.Frame(chat_frame)
        input_frame.pack(fill=tk.X, padx=10, pady=5)
        
        # User input
        self.user_input = tk.Text(input_frame, height=3, font=('Arial', 11))
        self.user_input.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 5))
        self.user_input.bind('<Control-Return>', self.send_message)
        
        # Buttons
        button_frame = ttk.Frame(input_frame)
        button_frame.pack(side=tk.RIGHT, fill=tk.Y)
        
        send_btn = ttk.Button(button_frame, text="Send (Ctrl+Enter)", command=self.send_message)
        send_btn.pack(fill=tk.X, pady=(0, 5))
        
        clear_btn = ttk.Button(button_frame, text="Clear Chat", command=self.clear_chat)
        clear_btn.pack(fill=tk.X)
        
        # Status frame
        status_frame = ttk.Frame(chat_frame)
        status_frame.pack(fill=tk.X, padx=10, pady=5)
        
        self.status_label = ttk.Label(status_frame, text="Ready - AI Models Loaded")
        self.status_label.pack(side=tk.LEFT)
        
        self.processing_label = ttk.Label(status_frame, text="")
        self.processing_label.pack(side=tk.RIGHT)
        
    def setup_database_tab(self):
        """Setup database management interface"""
        db_frame = ttk.Frame(self.notebook)
        self.notebook.add(db_frame, text="🗄️ Database")
        
        # Product management
        product_frame = ttk.LabelFrame(db_frame, text="Product Database")
        product_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # Buttons
        btn_frame = ttk.Frame(product_frame)
        btn_frame.pack(fill=tk.X, padx=5, pady=5)
        
        ttk.Button(btn_frame, text="Add Sample Products", 
                  command=self.add_sample_products).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="View Products", 
                  command=self.view_products).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Search Test", 
                  command=self.test_search).pack(side=tk.LEFT, padx=5)
        
        # Product display
        self.product_display = scrolledtext.ScrolledText(
            product_frame, height=15, state=tk.DISABLED
        )
        self.product_display.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
    def setup_training_tab(self):
        """Setup training interface"""
        training_frame = ttk.Frame(self.notebook)
        self.notebook.add(training_frame, text="📚 Training")
        
        # File upload section
        upload_frame = ttk.LabelFrame(training_frame, text="Upload Training Data")
        upload_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Button(upload_frame, text="Upload Excel Files", 
                  command=lambda: self.process_files('excel')).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(upload_frame, text="Upload PDF Files", 
                  command=lambda: self.process_files('pdf')).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(upload_frame, text="Upload Word Files", 
                  command=lambda: self.process_files('word')).pack(side=tk.LEFT, padx=5, pady=5)
        
        # Training status
        self.training_status = scrolledtext.ScrolledText(
            training_frame, height=20, state=tk.DISABLED
        )
        self.training_status.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
    def setup_analytics_tab(self):
        """Setup analytics interface"""
        analytics_frame = ttk.Frame(self.notebook)
        self.notebook.add(analytics_frame, text="📊 Analytics")
        
        # Metrics frame
        metrics_frame = ttk.LabelFrame(analytics_frame, text="Performance Metrics")
        metrics_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Button(metrics_frame, text="Refresh Analytics", 
                  command=self.refresh_analytics).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(metrics_frame, text="Export Data", 
                  command=self.export_analytics).pack(side=tk.LEFT, padx=5, pady=5)
        
        # Analytics display
        self.analytics_display = scrolledtext.ScrolledText(
            analytics_frame, height=25, state=tk.DISABLED
        )
        self.analytics_display.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
    def send_message(self, event=None):
        """Handle sending user message"""
        user_text = self.user_input.get(1.0, tk.END).strip()
        if not user_text:
            return
            
        # Clear input
        self.user_input.delete(1.0, tk.END)
        
        # Display user message
        self.display_message("You", user_text, "user")
        
        # Update status
        self.processing_label.config(text="Processing...")
        self.root.update()
        
        # Process in separate thread
        threading.Thread(
            target=self.process_message_thread,
            args=(user_text,),
            daemon=True
        ).start()
        
    def process_message_thread(self, user_input):
        """Process message in separate thread"""
        try:
            result = self.process_user_message(user_input)
            
            # Display response
            self.display_message("AI Assistant", result['response'], "bot")
            
            # Update status
            status_text = f"Response from: {result['data_source']} | "
            status_text += f"Local results: {result['local_results_count']} | "
            status_text += f"Time: {result['processing_time']:.2f}s"
            
            self.root.after(0, lambda: self.processing_label.config(text=status_text))
            
        except Exception as e:
            error_msg = f"Error processing message: {str(e)}"
            self.display_message("System", error_msg, "system")
            self.root.after(0, lambda: self.processing_label.config(text="Error"))
            
    def display_message(self, sender, message, tag):
        """Display message in chat window"""
        def update_display():
            self.chat_display.config(state=tk.NORMAL)
            timestamp = datetime.now().strftime("%H:%M:%S")
            
            self.chat_display.insert(tk.END, f"[{timestamp}] {sender}:\n", tag)
            self.chat_display.insert(tk.END, f"{message}\n\n")
            
            self.chat_display.config(state=tk.DISABLED)
            self.chat_display.see(tk.END)
            
        if threading.current_thread() != threading.main_thread():
            self.root.after(0, update_display)
        else:
            update_display()
            
    def clear_chat(self):
        """Clear chat display"""
        self.chat_display.config(state=tk.NORMAL)
        self.chat_display.delete(1.0, tk.END)
        self.chat_display.config(state=tk.DISABLED)
        self.conversation_context.clear()
        
    def add_sample_products(self):
        """Add sample products to database"""
        sample_products = [
            {
                'name': 'Gaming Laptop Pro X1',
                'description': 'High-performance gaming laptop with RTX 4070, Intel i7, 32GB RAM',
                'category': 'Laptops',
                'price': 1899.99,
                'features': 'RTX 4070, Intel i7-12700H, 32GB DDR5, 1TB NVMe SSD',
                'specifications': '15.6" 144Hz display, Windows 11, RGB keyboard',
                'availability': 'In Stock'
            },
            {
                'name': 'Business Ultrabook Z5',
                'description': 'Lightweight business laptop with long battery life',
                'category': 'Laptops',
                'price': 1299.99,
                'features': 'Intel i5-12500U, 16GB RAM, 512GB SSD, 14-inch',
                'specifications': '14" FHD display, 12-hour battery, Windows 11 Pro',
                'availability': 'In Stock'
            },
            {
                'name': 'Wireless Gaming Mouse RGB',
                'description': 'High-precision wireless gaming mouse with customizable RGB',
                'category': 'Accessories',
                'price': 79.99,
                'features': '16000 DPI, wireless charging, 7 programmable buttons',
                'specifications': '2.4GHz wireless, USB-C charging, 70-hour battery',
                'availability': 'In Stock'
            }
        ]
        
        try:
            cursor = self.conn.cursor()
            for product in sample_products:
                cursor.execute('''
                    INSERT INTO products 
                    (name, description, category, price, features, specifications, availability, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    product['name'],
                    product['description'],
                    product['category'],
                    product['price'],
                    product['features'],
                    product['specifications'],
                    product['availability'],
                    datetime.now().isoformat()
                ))
            
            self.conn.commit()
            messagebox.showinfo("Success", f"Added {len(sample_products)} sample products to database!")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error adding sample products: {e}")
            
    def view_products(self):
        """View products in database"""
        try:
            cursor = self.conn.cursor()
            cursor.execute("SELECT name, description, category, price, availability FROM products")
            products = cursor.fetchall()
            
            self.product_display.config(state=tk.NORMAL)
            self.product_display.delete(1.0, tk.END)
            
            if products:
                self.product_display.insert(tk.END, f"Found {len(products)} products in database:\n\n")
                
                for i, product in enumerate(products, 1):
                    self.product_display.insert(tk.END, f"{i}. {product[0]}\n")
                    self.product_display.insert(tk.END, f"   Description: {product[1][:100]}...\n")
                    self.product_display.insert(tk.END, f"   Category: {product[2]}\n")
                    self.product_display.insert(tk.END, f"   Price: ${product[3]}\n")
                    self.product_display.insert(tk.END, f"   Availability: {product[4]}\n\n")
            else:
                self.product_display.insert(tk.END, "No products found in database.\n")
                self.product_display.insert(tk.END, "Click 'Add Sample Products' to add test data.")
                
            self.product_display.config(state=tk.DISABLED)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error viewing products: {e}")
            
    def test_search(self):
        """Test search functionality"""
        test_queries = [
            "gaming laptop",
            "business computer",
            "wireless mouse",
            "cheap laptop under $1500",
            "RGB accessories"
        ]
        
        self.product_display.config(state=tk.NORMAL)
        self.product_display.delete(1.0, tk.END)
        self.product_display.insert(tk.END, "Testing search functionality...\n\n")
        
        for query in test_queries:
            self.product_display.insert(tk.END, f"Query: '{query}'\n")
            results = self.search_local_database(query)
            
            if results:
                for result in results[:2]:  # Show top 2
                    self.product_display.insert(tk.END, f"  ✓ {result['name']} (similarity: {result['similarity']:.3f})\n")
            else:
                self.product_display.insert(tk.END, "  ✗ No results found\n")
            
            self.product_display.insert(tk.END, "\n")
            
        self.product_display.config(state=tk.DISABLED)
        
    def process_files(self, file_type):
        """Process different types of files for training data"""
        file_types = {
            'excel': [('Excel files', '*.xlsx *.xls *.csv')],
            'pdf': [('PDF files', '*.pdf')],
            'word': [('Word files', '*.docx *.doc')]
        }
        
        files = filedialog.askopenfilenames(
            title=f"Select {file_type} files",
            filetypes=file_types[file_type]
        )
        
        if files:
            threading.Thread(
                target=self.process_files_worker,
                args=(files, file_type),
                daemon=True
            ).start()
            
    def process_files_worker(self, files, file_type):
        """Worker thread for processing files"""
        for file_path in files:
            try:
                self.update_training_status(f"Processing: {os.path.basename(file_path)}")
                
                if file_type == 'excel':
                    self.process_excel_file(file_path)
                elif file_type == 'pdf':
                    self.process_pdf_file(file_path)
                elif file_type == 'word':
                    self.process_word_file(file_path)
                    
                self.update_training_status(f"✓ Completed: {os.path.basename(file_path)}")
                
            except Exception as e:
                self.update_training_status(f"✗ Error processing {os.path.basename(file_path)}: {e}")
                
    def process_excel_file(self, file_path):
        """Process Excel files for product data"""
        try:
            # Read Excel file
            if file_path.endswith('.csv'):
                df = pd.read_csv(file_path)
            else:
                df = pd.read_excel(file_path)
                
            # Expected columns: name, description, category, price, features, specifications
            required_columns = ['name', 'description']
            
            if not all(col in df.columns for col in required_columns):
                raise ValueError(f"Excel file must contain columns: {required_columns}")
                
            cursor = self.conn.cursor()
            
            for _, row in df.iterrows():
                cursor.execute('''
                    INSERT INTO products 
                    (name, description, category, price, features, specifications, availability, source_file, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    str(row.get('name', '')),
                    str(row.get('description', '')),
                    str(row.get('category', 'General')),
                    float(row.get('price', 0)) if pd.notna(row.get('price')) else 0,
                    str(row.get('features', '')),
                    str(row.get('specifications', '')),
                    str(row.get('availability', 'Unknown')),
                    file_path,
                    datetime.now().isoformat()
                ))
                
            self.conn.commit()
            self.update_training_status(f"Added {len(df)} products from Excel file")
            
        except Exception as e:
            raise Exception(f"Error processing Excel file: {e}")
            
    def process_pdf_file(self, file_path):
        """Process PDF files for knowledge base"""
        try:
            reader = PdfReader(file_path)
            text = ""
            
            for page in reader.pages:
                text += page.extract_text() + "\n"
                
            if text.strip():
                # Store in knowledge base
                cursor = self.conn.cursor()
                cursor.execute('''
                    INSERT INTO knowledge_base 
                    (topic, content, source, created_at)
                    VALUES (?, ?, ?, ?)
                ''', (
                    os.path.basename(file_path),
                    text[:5000],  # Limit to 5000 chars
                    file_path,
                    datetime.now().isoformat()
                ))
                
                self.conn.commit()
                self.update_training_status(f"Added knowledge from PDF: {len(text)} characters")
                
        except Exception as e:
            raise Exception(f"Error processing PDF file: {e}")
            
    def process_word_file(self, file_path):
        """Process Word documents for knowledge base"""
        try:
            doc = docx.Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
                
            if text.strip():
                # Store in knowledge base
                cursor = self.conn.cursor()
                cursor.execute('''
                    INSERT INTO knowledge_base 
                    (topic, content, source, created_at)
                    VALUES (?, ?, ?, ?)
                ''', (
                    os.path.basename(file_path),
                    text[:5000],  # Limit to 5000 chars
                    file_path,
                    datetime.now().isoformat()
                ))
                
                self.conn.commit()
                self.update_training_status(f"Added knowledge from Word doc: {len(text)} characters")
                
        except Exception as e:
            raise Exception(f"Error processing Word file: {e}")
            
    def update_training_status(self, message):
        """Update training status display"""
        def update():
            self.training_status.config(state=tk.NORMAL)
            self.training_status.insert(tk.END, f"[{datetime.now().strftime('%H:%M:%S')}] {message}\n")
            self.training_status.config(state=tk.DISABLED)
            self.training_status.see(tk.END)
            
        if threading.current_thread() != threading.main_thread():
            self.root.after(0, update)
        else:
            update()
            
    def refresh_analytics(self):
        """Refresh analytics display"""
        try:
            cursor = self.conn.cursor()
            
            # Get conversation statistics
            cursor.execute("""
                SELECT 
                    COUNT(*) as total_conversations,
                    AVG(response_time) as avg_response_time,
                    data_source,
                    COUNT(*) as count_by_source
                FROM conversations 
                GROUP BY data_source
            """)
            
            source_stats = cursor.fetchall()
            
            # Get recent conversations
            cursor.execute("""
                SELECT timestamp, user_input, data_source, response_time
                FROM conversations 
                ORDER BY timestamp DESC 
                LIMIT 10
            """)
            
            recent_convs = cursor.fetchall()
            
            # Get product count
            cursor.execute("SELECT COUNT(*) FROM products")
            product_count = cursor.fetchone()[0]
            
            # Display analytics
            self.analytics_display.config(state=tk.NORMAL)
            self.analytics_display.delete(1.0, tk.END)
            
            # Summary
            self.analytics_display.insert(tk.END, "=== CHATBOT ANALYTICS ===\n\n")
            self.analytics_display.insert(tk.END, f"Products in Database: {product_count}\n")
            
            # Data source statistics
            self.analytics_display.insert(tk.END, "\n=== Response Sources ===\n")
            for stat in source_stats:
                self.analytics_display.insert(tk.END, f"{stat[2]}: {stat[3]} responses\n")
                
            # Performance metrics
            if source_stats:
                total_conversations = sum(stat[3] for stat in source_stats)
                avg_response_time = sum(stat[1] * stat[3] for stat in source_stats if stat[1]) / total_conversations
                
                self.analytics_display.insert(tk.END, f"\nTotal Conversations: {total_conversations}\n")
                self.analytics_display.insert(tk.END, f"Average Response Time: {avg_response_time:.2f} seconds\n")
                
            # Recent conversations
            self.analytics_display.insert(tk.END, "\n=== Recent Conversations ===\n")
            for conv in recent_convs:
                timestamp, user_input, data_source, response_time = conv
                self.analytics_display.insert(tk.END, f"{timestamp} - {data_source} ({response_time:.2f}s)\n")
                self.analytics_display.insert(tk.END, f"User: {user_input[:100]}...\n\n")
                
            self.analytics_display.config(state=tk.DISABLED)
            
        except Exception as e:
            messagebox.showerror("Error", f"Error refreshing analytics: {e}")
            
    def export_analytics(self):
        """Export analytics data to CSV"""
        try:
            cursor = self.conn.cursor()
            cursor.execute("SELECT * FROM conversations")
            
            df = pd.DataFrame(
                cursor.fetchall(),
                columns=['id', 'timestamp', 'user_input', 'bot_response', 'intent', 'confidence', 'data_source', 'response_time']
            )
            
            export_path = f"chatbot_analytics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
            df.to_csv(export_path, index=False)
            
            messagebox.showinfo("Success", f"Analytics exported to: {export_path}")
            
        except Exception as e:
            messagebox.showerror("Error", f"Error exporting analytics: {e}")
            
    def run(self):
        """Start the application"""
        # Display welcome message
        welcome_msg = """🤖 Natural Language AI Sales ChatBot Initialized!

Features:
✅ Natural language text generation using local LLM
✅ Local database search with semantic similarity
✅ Google search fallback for unknown queries
✅ Conversation context awareness
✅ Multi-GPU support (RTX 4090 + RTX 4070Ti Super)

Try asking questions like:
• "Tell me about gaming laptops under $2000"
• "What wireless accessories do you have?"
• "I need a business computer for presentations"
• "Compare your laptop models"

The AI will search the local database first, then Google if needed, and generate natural responses using the LLM."""

        self.display_message("System", welcome_msg, "system")
        
        # Start the GUI
        self.root.mainloop()
        
    def __del__(self):
        """Cleanup"""
        if hasattr(self, 'conn'):
            self.conn.close()

# Example usage and testing functions
def create_enhanced_chatbot():
    """Create and return an enhanced chatbot instance"""
    return NaturalLanguageSalesChatBot()

def test_natural_language_responses():
    """Test the natural language generation capabilities"""
    print("Testing Natural Language AI ChatBot...")
    
    # Test queries
    test_queries = [
        "I'm looking for a gaming laptop under $2000",
        "What business computers do you recommend?",
        "Tell me about wireless gaming accessories",
        "What's the difference between your laptop models?",
        "I need help choosing between Intel and AMD processors"
    ]
    
    chatbot = create_enhanced_chatbot()
    
    for query in test_queries:
        print(f"\n🔍 Query: {query}")
        result = chatbot.process_user_message(query)
        print(f"📊 Source: {result['data_source']}")
        print(f"⏱️ Time: {result['processing_time']:.2f}s")
        print(f"🤖 Response: {result['response'][:200]}...")

if __name__ == "__main__":
    # Check for required directories
    os.makedirs('data', exist_ok=True)
    os.makedirs('logs', exist_ok=True)
    
    print("""
    ===============================================
    🚀 Natural Language AI Sales ChatBot
    ===============================================
    
    Initializing AI models...
    This may take a few minutes on first run.
    """)
    
    try:
        # Create and run the chatbot
        chatbot = NaturalLanguageSalesChatBot()
        chatbot.run()
        
    except KeyboardInterrupt:
        print("\n👋 ChatBot stopped by user")
    except Exception as e:
        print(f"❌ Error starting chatbot: {e}")
        input("Press Enter to exit...")